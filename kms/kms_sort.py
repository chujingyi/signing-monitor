"""生成 KMStation 活动的库存变动排序报告（Excel 格式）。

日志行格式：
  [2026-05-29 10:00:00] prodId=4013 skuId=123 sion 100 → 95 (Δ-5)
"""

from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# 匹配：[ts] prodId=X skuId=Y {sku_name} {old} → {new} (Δ{diff})
LOG_RE = re.compile(
    r'^\[(.+?)\] prodId=\S+ skuId=\S+ (.+?) \d+ → \d+ \(Δ([+-]\d+)\)'
)


def _parse_log(stock_log: Path) -> dict[str, list[dict]]:
    """解析 stock.log，返回 {sku_name: [{ts, sku_name, delta}]}"""
    member_records: dict[str, list[dict]] = defaultdict(list)
    with stock_log.open(encoding="utf-8") as f:
        for line in f:
            m = LOG_RE.match(line.strip())
            if not m:
                continue
            ts       = m.group(1)
            sku_name = m.group(2).strip()
            delta    = int(m.group(3))
            member_records[sku_name].append({"ts": ts, "sku_name": sku_name, "delta": delta})
    return member_records


def _parse_name_parts(name: str) -> tuple[str, str]:
    """从活动名称提取 (活动类型, 艺人名称)。
    name 格式示例：【个人签售】NCT WISH - 1st Album...
    """
    activity_type = ""
    artist = ""
    m = re.search(r"【(.+?)】", name)
    if m:
        activity_type = m.group(1)
        rest = name[m.end():].strip()
    else:
        rest = name.strip()
    # 艺人名：截到第一个 " - " / "- " / "[" / 中文字符
    m2 = re.search(r"\s*[-–]\s|\[|[一-鿿]", rest)
    if m2:
        artist = rest[:m2.start()].strip()
    else:
        artist = rest[:20].strip()
    return activity_type, artist


def _build_xlsx(member_records: dict[str, list[dict]], name: str, exports_dir: Path) -> Path:
    """生成 Excel 文件，每个成员一个 sheet，返回文件路径。"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # 删除默认空 sheet

    # 按累计变化从小到大排序成员（最多应募排最前）
    member_total  = {m: sum(r["delta"] for r in recs) for m, recs in member_records.items()}
    sorted_members = sorted(member_total, key=lambda m: member_total[m])

    header_font  = Font(bold=True, color="FFFFFF")
    header_fill  = PatternFill("solid", fgColor="4472C4")
    center_align = Alignment(horizontal="center")

    for member in sorted_members:
        recs = member_records[member]
        total = member_total[member]
        # Sheet 名不能超过 31 字符
        sheet_name = member[:31]
        ws = wb.create_sheet(title=sheet_name)

        # 标题行
        ws.append(["时间", "SKU名称", "变化量"])
        for cell in ws[1]:
            cell.font  = header_font
            cell.fill  = header_fill
            cell.alignment = center_align

        # 数据行：按变化量从小到大排序（最负排最前）
        for r in sorted(recs, key=lambda x: x["delta"]):
            ws.append([r["ts"], r["sku_name"], r["delta"]])

        # 汇总行
        ws.append([])
        ws.append(["", "累计变化", total])
        summary_row = ws.max_row
        for cell in ws[summary_row]:
            cell.font = Font(bold=True)

        # 列宽
        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 40
        ws.column_dimensions["C"].width = 10

    dt_tag = datetime.now().strftime("%Y%m%d_%H%M")
    activity_type, artist = _parse_name_parts(name)
    # 文件名安全处理：空格→下划线，去掉其他特殊符号
    def _safe(s: str) -> str:
        return re.sub(r"[^\w一-鿿]+", "_", s).strip("_")
    fname = f"{dt_tag}_{_safe(artist)}_{_safe(activity_type)}.xlsx" if activity_type or artist else f"{dt_tag}_rank.xlsx"
    # 删除旧的 xlsx，保持目录里只有最新一个
    for old_xlsx in exports_dir.glob("*.xlsx"):
        old_xlsx.unlink()
    export_path = exports_dir / fname
    wb.save(export_path)
    return export_path


def sort_campaign(config_path: Path) -> tuple[str, Path | None]:
    """返回 (飞书文字摘要, xlsx路径)；无数据时 xlsx路径为 None。"""
    config_path = Path(config_path).resolve()
    with config_path.open(encoding="utf-8-sig") as f:
        cfg = json.load(f)

    campaign_dir = config_path.parent
    stock_log    = campaign_dir / "logs" / "stock.log"
    exports_dir  = campaign_dir / "exports"
    exports_dir.mkdir(exist_ok=True)

    name = cfg.get("name") or cfg.get("campaign_id") or campaign_dir.name

    if not stock_log.exists():
        return f"KMS 报告：{name}\n\n暂无变化记录", None

    member_records = _parse_log(stock_log)
    total = sum(len(v) for v in member_records.values())

    if total == 0:
        return f"KMS 报告：{name}\n\n暂无变化记录", None

    # 生成 Excel
    xlsx_path = _build_xlsx(member_records, name, exports_dir)

    # 飞书文字摘要
    member_total   = {m: sum(r["delta"] for r in recs) for m, recs in member_records.items()}
    sorted_members = sorted(member_total, key=lambda m: member_total[m])

    lines = [f"KMS 排序报告：{name}", f"{total} 条变化记录\n"]
    for member in sorted_members:
        recs        = member_records[member]
        total_delta = member_total[member]
        lines.append(f"【{member}】 {total_delta:+d}（{len(recs)} 次）")
        for r in sorted(recs, key=lambda x: x["delta"]):
            lines.append(f"  {r['ts']}  {r['delta']:+d}")

    return "\n".join(lines).rstrip(), xlsx_path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", required=True)
    args = parser.parse_args()
    text, xlsx = sort_campaign(Path(args.config))
    print(text)
    if xlsx:
        print(f"\nExcel 已生成: {xlsx}")


if __name__ == "__main__":
    main()
